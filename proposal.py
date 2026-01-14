from docx import Document
from datetime import date

def brl(v: float) -> str:
    s = f"{v:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X",".")

def gerar_proposta_docx(out: dict, cliente: str, cnpj: str = "", responsavel: str = "") -> str:
    doc = Document()
    doc.add_heading('Proposta Comercial — Consultoria Tributária', level=1)
    doc.add_paragraph(f"Data: {date.today().strftime('%d/%m/%Y')}")
    if responsavel:
        doc.add_paragraph(f"Responsável: {responsavel}")

    doc.add_heading("1. Identificação do Cliente", level=2)
    doc.add_paragraph(f"Cliente: {cliente}")
    if cnpj:
        doc.add_paragraph(f"CNPJ: {cnpj}")

    doc.add_heading("2. Objeto e Escopo", level=2)
    doc.add_paragraph(out["escopo_para_proposta"])

    doc.add_heading("3. Prazo", level=2)
    doc.add_paragraph(f"Prazo estimado: {out.get('prazo_estimado_dias','—')} dias.")

    doc.add_heading("4. Equipe e Alocação", level=2)
    for cargo, horas in out["horas_por_cargo"].items():
        doc.add_paragraph(f"- {cargo}: {horas} horas")

    doc.add_heading("5. Investimento", level=2)
    doc.add_paragraph(f"Custo interno estimado: R$ {brl(out['custo_entrega'])}")
    doc.add_paragraph(f"Preço de venda: R$ {brl(out['preco_venda'])}")

    doc.add_heading("6. Condições Gerais", level=2)
    doc.add_paragraph("- Validade: 10 dias.")
    doc.add_paragraph("- Início após aceite e kickoff.")
    doc.add_paragraph("- Premissas/exclusões conforme escopo.")

    filename = f"Proposta_{cliente.replace(' ', '_')}.docx"
    doc.save(filename)
    return filename
